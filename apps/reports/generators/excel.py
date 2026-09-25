from __future__ import annotations

import hashlib
import os
from datetime import date, datetime
from tempfile import NamedTemporaryFile
from pathlib import Path

from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from apps.companies.models import State
from apps.customers.visibility import exclude_hidden_customers
from apps.processing.services import (
    calculate_indicators,
    calculate_overdue_bands,
    calculate_overdue_days,
    classify_overdue_days,
)
from apps.receivables.models import Receivable
from apps.reports.models import ReportExecution
from apps.reports.retention import supersede_old_reports

SHEETS = ("5_10", "11_30", "31_90", "91_360")
BAND_LABELS = {
    "current": "Em dia ou até 4 dias",
    "5_10": "De 5 a 10 dias",
    "11_30": "De 11 a 30 dias",
    "31_90": "De 31 a 90 dias",
    "91_360": "De 91 a 360 dias",
    "over_360": "Acima de 360 dias",
}
BAND_FILLS = {
    "5_10": "0CA30C",
    "11_30": "D9A300",
    "31_90": "E87511",
    "91_360": "EF8A8A",
}
BAND_FILENAME_LABELS = {
    "5_10": "5-10",
    "11_30": "11-30",
    "31_90": "31-90",
    "91_360": "91-360",
}
TITLE_HEADERS = [
    "ID do título", "Parcela", "ID do cliente", "Grupo",
    "Empresa", "Estado", "Valor original", "Juros", "Multa", "Saldo em aberto",
    "Emissão", "Vencimento", "Dias vencidos", "Faixa", "Status financeiro",
    "Juros ao dia", "Status na origem", "Vendedor", "Código do agente",
    "Estabelecimento de origem", "Código do estabelecimento",
    "Gestor responsável", "Status da cobrança", "Previsão de pagamento", "Última interação",
]


def _reports_dir() -> Path:
    target = Path(os.getenv("REPORTS_DIR", "reports"))
    target.mkdir(parents=True, exist_ok=True)
    return target


def _state_items(state: State) -> list[Receivable]:
    return list(
        exclude_hidden_customers(Receivable.objects.filter(company__state=state))
        .select_related("company__state", "customer__economic_group", "operator", "collection_status")
        .prefetch_related("interactions")
        .order_by("due_date", "id")
    )


def _title_row(item: Receivable, days: int, band_label: str) -> list[object]:
    collection = getattr(item, "collection_status", None)
    last_interaction = item.interactions.first()
    responsible = ""
    if item.operator:
        responsible = item.operator.get_full_name() or item.operator.username
    return [
        item.title_number, item.installment, item.customer.identifier,
        item.customer.economic_group.name if item.customer.economic_group else "",
        item.company.name, item.company.state.code, item.original_amount, item.interest_amount,
        item.penalty_amount, item.outstanding_amount, item.issued_at, item.due_date,
        days, band_label,
        item.get_financial_status_display(), item.daily_interest_amount, item.source_status,
        item.seller, item.agent_code, item.origin_establishment_code, item.establishment_code, responsible,
        collection.get_status_display() if collection else "Pendente",
        collection.expected_payment_date if collection else None,
        last_interaction.created_at.replace(tzinfo=None) if last_interaction else None,
    ]


def _format_sheet(sheet) -> None:
    header_fill = PatternFill("solid", fgColor="6E4B2E")
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center")
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for row in sheet.iter_rows():
        for cell in row:
            if isinstance(cell.value, datetime):
                cell.number_format = "dd/mm/yyyy hh:mm"
            elif isinstance(cell.value, date):
                cell.number_format = "dd/mm/yyyy"
    for index, column in enumerate(sheet.columns, 1):
        width = min(max(len(str(cell.value or "")) for cell in column) + 2, 36)
        sheet.column_dimensions[get_column_letter(index)].width = width


def _format_summary(sheet) -> None:
    _format_sheet(sheet)
    sheet.column_dimensions["A"].width = 32
    sheet.column_dimensions["B"].width = 22
    row_fills = {
        6: "6E4B2E",
        7: BAND_FILLS["5_10"],
        8: BAND_FILLS["11_30"],
        9: BAND_FILLS["31_90"],
        10: BAND_FILLS["91_360"],
    }
    for row, color in row_fills.items():
        label = sheet.cell(row=row, column=1)
        label.fill = PatternFill("solid", fgColor=color)
        font_color = "FFFFFF" if row in {6, 7, 9} else "1F1B17"
        label.font = Font(color=font_color, bold=True)
    for row in range(5, 11):
        sheet.cell(row=row, column=2).number_format = '#,##0.00'
    sheet.cell(row=11, column=2).number_format = '0.00'
    sheet.cell(row=12, column=2).number_format = '#,##0'
    for row in range(2, sheet.max_row + 1):
        sheet.cell(row=row, column=2).alignment = Alignment(horizontal="right")


def _save_execution(execution: ReportExecution, workbook: Workbook, path: Path, rows: int, metadata: dict) -> None:
    with NamedTemporaryFile(dir=path.parent, suffix=path.suffix, delete=False) as temporary:
        temporary_path = Path(temporary.name)
    try:
        workbook.save(temporary_path)
        checksum = hashlib.sha256(temporary_path.read_bytes()).hexdigest()
        temporary_path.replace(path)
    finally:
        temporary_path.unlink(missing_ok=True)
    execution.file_path = str(path.resolve())
    execution.checksum = checksum
    execution.rows_processed = rows
    execution.status = ReportExecution.Status.COMPLETED
    execution.metadata = metadata


def generate_state_excel(state_code: str, reference_date: date) -> ReportExecution:
    state, _ = State.objects.get_or_create(code=state_code.upper(), defaults={"name": state_code.upper()})
    execution, _ = ReportExecution.objects.get_or_create(
        state=state, report_type=ReportExecution.Type.STATE_EXCEL,
        overdue_band="", report_date=reference_date,
    )
    execution.status = ReportExecution.Status.PROCESSING
    execution.started_at = timezone.now()
    execution.error_message = ""
    execution.save()
    try:
        items = _state_items(state)
        indicators = calculate_indicators(items, reference_date)
        overdue_bands = calculate_overdue_bands(items, reference_date)
        workbook = Workbook()
        summary = workbook.active
        summary.title = "Resumo"
        summary.append(["Indicador", "Valor"])
        summary.append(["Estado", state.code])
        summary.append(["Empresas", "Nova + Multi" if state.code == "PE" else "Todas do estado"])
        summary.append(["Data de referência", reference_date])
        summary.append(["Carteira total", indicators["portfolio_total"]])
        summary.append(["Vencidos total", overdue_bands["overdue_total"]])
        summary.append(["Vencidos 5-10 dias", overdue_bands["days_5_10"]])
        summary.append(["Vencidos 11-30 dias", overdue_bands["days_11_30"]])
        summary.append(["Vencidos 31-90 dias", overdue_bands["days_31_90"]])
        summary.append(["Vencidos 91-360 dias", overdue_bands["days_91_360"]])
        summary.append(["Inadimplência (%)", indicators["delinquency_percentage"]])
        summary.append(["Quantidade de títulos", indicators["title_count"]])
        _format_summary(summary)

        all_titles = workbook.create_sheet("Todas as duplicatas")
        all_titles.append(TITLE_HEADERS)
        sheets = {name: workbook.create_sheet(name) for name in SHEETS}
        for name, sheet in sheets.items():
            sheet.sheet_properties.tabColor = BAND_FILLS[name]
            sheet.append(TITLE_HEADERS)
        written = 0
        for item in items:
            days = calculate_overdue_days(item.due_date, reference_date)
            band = classify_overdue_days(days).value
            all_titles.append(_title_row(item, days, BAND_LABELS[band]))
            if band in sheets:
                sheet = sheets[band]
                sheet.append(_title_row(item, days, BAND_LABELS[band]))
                written += 1
        _format_sheet(all_titles)
        for sheet in sheets.values():
            _format_sheet(sheet)
        band_filename = "_".join(BAND_FILENAME_LABELS[name] for name in SHEETS)
        path = _reports_dir() / f"cobranca_{state.code}_{band_filename}.xlsx"
        _save_execution(execution, workbook, path, written, {
            "sheets": ["Resumo", "Todas as duplicatas", *SHEETS],
        })
    except Exception as exc:
        execution.status = ReportExecution.Status.FAILED
        execution.error_message = str(exc)
        raise
    finally:
        execution.finished_at = timezone.now()
        execution.save()
    supersede_old_reports(execution, _reports_dir())
    return execution


def generate_upcoming_excel(state_code: str, reference_date: date) -> ReportExecution:
    state, _ = State.objects.get_or_create(code=state_code.upper(), defaults={"name": state_code.upper()})
    execution, _ = ReportExecution.objects.get_or_create(
        state=state, report_type=ReportExecution.Type.UPCOMING_EXCEL,
        overdue_band="", report_date=reference_date,
    )
    execution.status = ReportExecution.Status.PROCESSING
    execution.started_at = timezone.now()
    execution.error_message = ""
    execution.save()
    try:
        items = [item for item in _state_items(state) if item.due_date >= reference_date]
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "A vencer"
        sheet.append(TITLE_HEADERS)
        for item in items:
            sheet.append(_title_row(item, 0, "A vencer"))
        _format_sheet(sheet)
        path = _reports_dir() / f"cobranca_a_vencer_{state.code}.xlsx"
        _save_execution(execution, workbook, path, len(items), {
            "companies": "Nova + Multi" if state.code == "PE" else "Todas do estado",
        })
    except Exception as exc:
        execution.status = ReportExecution.Status.FAILED
        execution.error_message = str(exc)
        raise
    finally:
        execution.finished_at = timezone.now()
        execution.save()
    supersede_old_reports(execution, _reports_dir())
    return execution
