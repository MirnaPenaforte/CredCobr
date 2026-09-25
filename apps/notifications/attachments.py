from __future__ import annotations

import math
from copy import copy
from datetime import date, datetime
from pathlib import Path

from django.conf import settings
from openpyxl import Workbook, load_workbook

SUMMARY_SHEET = "Resumo"
DUPLICATES_SHEET = "Todas as duplicatas"
ALWAYS_SHEETS = (SUMMARY_SHEET, DUPLICATES_SHEET)
BAND_GROUPS = (
    ("5_10", "11_30"),
    ("31_90", "91_360"),
)
BAND_TAB_COLORS = {
    "5_10": "0CA30C",
    "11_30": "D9A300",
    "31_90": "E87511",
    "91_360": "EF8A8A",
}
BAND_FILENAME_LABELS = {
    "5_10": "5-10",
    "11_30": "11-30",
    "31_90": "31-90",
    "91_360": "91-360",
}


def _copy_cell_style(source, target) -> None:
    # Read-only worksheets expose style components but not the private
    # ``_style`` object, so copy the public components individually.
    if getattr(source, "has_style", False):
        target.font = copy(source.font)
        target.fill = copy(source.fill)
        target.border = copy(source.border)
        target.alignment = copy(source.alignment)
        target.protection = copy(source.protection)
        target.number_format = source.number_format


def _copy_formatting(sheet) -> None:
    for row in sheet.iter_rows():
        for cell in row:
            if isinstance(cell.value, datetime):
                cell.number_format = "dd/mm/yyyy hh:mm"
            elif isinstance(cell.value, date):
                cell.number_format = "dd/mm/yyyy"


def _append_row(source_row, target_sheet, *, copy_style: bool = False) -> None:
    target_sheet.append([cell.value for cell in source_row])
    if copy_style:
        target_row = target_sheet.max_row
        for source_cell, target_cell in zip(source_row, target_sheet[target_row]):
            _copy_cell_style(source_cell, target_cell)


def _workbook_sheetnames(source: Path) -> list[str]:
    workbook = load_workbook(source, read_only=True)
    names = list(workbook.sheetnames)
    workbook.close()
    return names


def _split_workbook(source: Path, parts: int) -> list[Path]:
    """Split an .xlsx report into ``parts`` files, splitting rows per sheet.

    The header row is kept in every part. The ``Resumo`` sheet stays whole in the
    first part. Rows are streamed, so large reports are not fully loaded in memory.
    """
    output: list[Path] = []
    for index in range(parts):
        workbook = load_workbook(source, read_only=True)
        part = Workbook()
        part.remove(part.active)
        for worksheet in workbook.worksheets:
            sheet = part.create_sheet(title=worksheet.title[:31])
            if worksheet.title in BAND_TAB_COLORS:
                sheet.sheet_properties.tabColor = BAND_TAB_COLORS[worksheet.title]
            data_total = max(worksheet.max_row - 1, 0)
            start = index * data_total // parts
            end = data_total if index == parts - 1 else (index + 1) * data_total // parts
            rows = worksheet.iter_rows()
            header = next(rows, None)
            if header is None:
                continue
            if worksheet.title == SUMMARY_SHEET:
                if index == 0:
                    _append_row(header, sheet, copy_style=True)
                    for row in rows:
                        _append_row(row, sheet, copy_style=True)
                continue
            _append_row(header, sheet, copy_style=True)
            for position, row in enumerate(rows):
                if start <= position < end:
                    _append_row(row, sheet)
            _copy_formatting(sheet)
        workbook.close()
        part_path = source.with_name(f"{source.stem}_parte{index + 1}{source.suffix}")
        part.save(part_path)
        output.append(part_path)
    return output


def _split_by_sheets(source: Path, groups: list[list[str]]) -> list[Path]:
    """Create one file per sheet group, copying the selected sheets."""
    output: list[Path] = []
    for index, sheet_names in enumerate(groups, start=1):
        workbook = load_workbook(source, read_only=True)
        part = Workbook()
        part.remove(part.active)
        for name in sheet_names:
            if name not in workbook.sheetnames:
                continue
            source_sheet = workbook[name]
            sheet = part.create_sheet(title=name[:31])
            if name in BAND_TAB_COLORS:
                sheet.sheet_properties.tabColor = BAND_TAB_COLORS[name]
            for row_index, row in enumerate(source_sheet.iter_rows()):
                _append_row(row, sheet, copy_style=row_index == 0 or name == SUMMARY_SHEET)
            _copy_formatting(sheet)
        workbook.close()
        part_bands = [name for name in BAND_FILENAME_LABELS if name in sheet_names]
        if part_bands:
            base = source.stem
            for label in BAND_FILENAME_LABELS.values():
                base = base.replace(f"_{label}", "")
            suffix = "_".join(BAND_FILENAME_LABELS[name] for name in part_bands)
            part_path = source.with_name(f"{base}_{suffix}{source.suffix}")
        else:
            part_path = source.with_name(f"{source.stem}_parte{index}{source.suffix}")
        part.save(part_path)
        output.append(part_path)
    return output


def _band_sheet_groups(sheet_names: list[str]) -> list[list[str]]:
    always = [name for name in ALWAYS_SHEETS if name in sheet_names]
    groups: list[list[str]] = []
    for band_group in BAND_GROUPS:
        bands = [name for name in band_group if name in sheet_names]
        if bands:
            groups.append([*always, *bands])
    return groups


def resolve_report_path(source: Path | str | None) -> Path | None:
    """Resolve a report path across host/container storage layouts."""
    if not source:
        return None
    path = Path(source)
    if path.is_file():
        return path

    # Older ReportExecution rows may contain the host's absolute path. Reports
    # shared by the web and worker containers live in the mounted reports dir.
    fallback = Path(settings.BASE_DIR) / "reports" / path.name
    return fallback if fallback.is_file() else None


def plan_report_attachments(source: Path | str | None, max_bytes: int | None = None) -> list[Path]:
    """Return the attachment plan for a report file.

    When the file fits the limit, a single-item list with the original file is
    returned. When it exceeds the limit, the report is split by overdue band:
    ``5_10``/``11_30`` go in one file and ``31_90``/``91_360`` in another, always
    carrying the ``Resumo`` and ``Todas as duplicatas`` sheets in both. Reports
    without band sheets fall back to a row split.
    """
    source = resolve_report_path(source)
    if source is None:
        return []
    limit = max_bytes or getattr(settings, "EMAIL_MAX_ATTACHMENT_BYTES", 18 * 1024 * 1024)
    size = source.stat().st_size
    if size <= limit:
        return [source]
    groups = _band_sheet_groups(_workbook_sheetnames(source))
    if len(groups) >= 2:
        return _split_by_sheets(source, groups)
    parts = max(2, math.ceil(size / limit))
    return _split_workbook(source, parts)
